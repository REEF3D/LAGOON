# SPDX-License-Identifier: GPL-3.0-or-later
# Architect: Hans Bihs
"""Data tables for 2D plots (no Qt): columns from files and from REEF3D's time series,
new columns from expressions, curve fits, and tables collected over several runs.

A :class:`TableSpec` says where a table comes from and which columns are computed,
so a figure keeps its data recipe and draws again when the files change or for
another run::

    spec = TableSpec("flume", source="file:measured/flume3.dat",
                     formulas={"eta_cm": "100 * P1", "H": "max(P1) - min(P1)"})
    table = spec.load(session, base=folder)       # a DataTable
    table.columns["eta_cm"]

Sources: ``file:<path>`` (text columns: spaces, tabs, commas or semicolons; a
header line of names; ``#`` and ``%`` comments; ``.xlsx`` with openpyxl),
``series:<key>`` (a time-series file of an open case: gauges, forces, motions),
``profile:<key>@<t>`` (a free surface line set at a time), ``collect:`` (one row
per open case, see :func:`collect`).

Expressions (:func:`evaluate`) use the columns by name (a name with spaces or
brackets as ``col("Fz [N]")``), numbers, ``+ - * / ** %``, comparisons, ``&``,
``|``, ``~`` and the functions in :data:`FUNCTIONS` (``sin``, ``sqrt``, ``mean``,
``rms``, ``hm0``, ``where``, ``clip``, ``cumsum``, ``diff``, ``gradient``,
``smooth``, ...). A result of one value fills the column.
"""

from __future__ import annotations

import ast
import math
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import numpy as np

_NUMBER = re.compile(r"^[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?$|^[-+]?(nan|inf)$", re.I)


# ======================================================================== tables
@dataclass
class DataTable:
    """Columns of equal length, in order, with units where known."""

    name: str
    columns: dict[str, np.ndarray] = field(default_factory=dict)
    units: dict[str, str] = field(default_factory=dict)
    source: str = ""
    #: the columns computed from expressions (name -> expression), in order
    formulas: dict[str, str] = field(default_factory=dict)
    #: problems computing them (name -> message)
    errors: dict[str, str] = field(default_factory=dict)
    #: names of the rows, when they have some (collect: the cases)
    labels: list[str] = field(default_factory=list)

    @property
    def rows(self) -> int:
        return max((len(v) for v in self.columns.values()), default=0)

    @property
    def names(self) -> list[str]:
        return list(self.columns)

    def column(self, name: str) -> np.ndarray:
        if name not in self.columns:
            raise KeyError(f"{self.name}: no column {name!r}")
        return self.columns[name]

    def add_formulas(self, formulas: dict[str, str]) -> None:
        """Compute columns from expressions, in order (a formula may use the ones
        before it); a failing one is left out and named in ``errors``."""
        for name, expression in formulas.items():
            self.formulas[name] = expression
            try:
                value = evaluate(expression, self.columns)
            except (ValueError, KeyError, TypeError, ZeroDivisionError, SyntaxError) as error:
                # (a KeyError's str() is its message in quotes)
                self.errors[name] = str(error.args[0]) if error.args else str(error)
                continue
            value = np.asarray(value, dtype=np.float64)
            if value.ndim == 0:
                value = np.full(self.rows, float(value))
            elif self.rows and value.size != self.rows:
                # diff and the like give one less: padded with NaN at the end
                padded = np.full(self.rows, np.nan)
                padded[: min(value.size, self.rows)] = value[: self.rows]
                value = padded
            self.columns[name] = value
            self.errors.pop(name, None)

    def text(self, limit: int | None = None) -> str:
        """Tab-separated, a header of names: for the clipboard or a file."""
        names = self.names
        lines = ["\t".join(names)]
        n = self.rows if limit is None else min(self.rows, limit)
        for k in range(n):
            lines.append(
                "\t".join(
                    f"{self.columns[c][k]:.10g}" if k < len(self.columns[c]) else "" for c in names
                )
            )
        return "\n".join(lines)


def read_table(path: str | Path, name: str = "") -> DataTable:
    """A table from a text file (columns separated by spaces, tabs, commas or
    semicolons; a header line of names; ``#`` / ``%`` comments; units in brackets
    in the names are kept as units) or an Excel file (.xlsx: the first sheet)."""
    path = Path(path).expanduser()
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_excel(path, name)
    names: list[str] | None = None
    rows: list[list[float]] = []
    comma_decimal = False
    for line in path.read_text(errors="replace").splitlines():
        text = line.strip()
        if not text or text[0] in "#%":
            continue
        if ";" in text and "," in text and not comma_decimal:
            comma_decimal = all(
                _NUMBER.match(f.replace(",", ".")) for f in re.split(r"[;\s]+", text) if f
            )
        separators = r"[;\s]+" if comma_decimal else r"[\s,;]+"
        fields = [f for f in re.split(separators, text) if f]
        if comma_decimal:
            fields = [f.replace(",", ".") for f in fields]
        if fields and all(_NUMBER.match(f) for f in fields):
            rows.append([float(f) for f in fields])
        elif not rows:
            names = _header(text)
    if not rows:
        raise ValueError(f"{path.name}: no numbers found")
    width = max(len(r) for r in rows)
    table = np.full((len(rows), width), np.nan)
    for k, row in enumerate(rows):
        table[k, : len(row)] = row
    return _from_matrix(name or path.stem, table, names, f"file:{path}")


def _header(text: str) -> list[str]:
    """Column names of a header line: tab or comma separated when there are tabs or
    commas (names may have spaces), else by spaces; "Fz [N]" keeps its unit."""
    if "\t" in text:
        parts = text.split("\t")
    elif "," in text or ";" in text:
        parts = re.split(r"[,;]", text)
    else:
        parts = re.findall(r"\S+(?:\s+\[[^\]]*\])?", text)
    return [p.strip() for p in parts if p.strip()]


def _from_matrix(name: str, table: np.ndarray, names: list[str] | None, source: str) -> DataTable:
    width = table.shape[1]
    if not names or len(names) < width:
        given = list(names or [])
        names = given + [f"c{k + 1}" for k in range(len(given), width)]
    out = DataTable(name, source=source)
    for k in range(width):
        label, unit = _unit(names[k])
        while label in out.columns:
            label += "_"
        out.columns[label] = table[:, k]
        if unit:
            out.units[label] = unit
    return out


def _unit(name: str) -> tuple[str, str]:
    match = re.match(r"^(.*?)\s*[\[(]([^\])]*)[\])]\s*$", name)
    if match and match.group(1):
        return match.group(1).strip(), match.group(2).strip()
    return name.strip(), ""


def _read_excel(path: Path, name: str) -> DataTable:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ValueError("reading .xlsx needs openpyxl (pip install openpyxl)") from error
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    names, rows = None, []
    for values in sheet.iter_rows(values_only=True):
        cells = [v for v in values]
        if all(v is None for v in cells):
            continue
        numeric = [isinstance(v, int | float) for v in cells if v is not None]
        if numeric and all(numeric):
            rows.append([float(v) if v is not None else math.nan for v in cells])
        elif not rows:
            names = [str(v) if v is not None else "" for v in cells]
    if not rows:
        raise ValueError(f"{path.name}: no numbers found")
    width = max(len(r) for r in rows)
    matrix = np.full((len(rows), width), np.nan)
    for k, row in enumerate(rows):
        matrix[k, : len(row)] = row
    return _from_matrix(name or path.stem, matrix, names, f"file:{path}")


# ======================================================================== expressions
def _rms(x):
    x = np.asarray(x, float)
    return float(np.sqrt(np.nanmean(x**2)))


def _hm0(x):
    x = np.asarray(x, float)
    return float(4.0 * np.nanstd(x))


def _smooth(x, n=5):
    """Running mean over ``n`` values (centred)."""
    x = np.asarray(x, float)
    n = max(int(n), 1)
    kernel = np.ones(n) / n
    return np.convolve(np.nan_to_num(x), kernel, mode="same")


def _interp(x, xp, fp):
    return np.interp(x, xp, fp)


def _range(x):
    x = np.asarray(x, float)
    return float(np.nanmax(x) - np.nanmin(x))


def _deg(x):
    return np.degrees(x)


def _rad(x):
    return np.radians(x)


#: what an expression may call
FUNCTIONS: dict[str, Any] = {
    # element by element
    "sin": np.sin,
    "cos": np.cos,
    "tan": np.tan,
    "arcsin": np.arcsin,
    "arccos": np.arccos,
    "arctan": np.arctan,
    "arctan2": np.arctan2,
    "sinh": np.sinh,
    "cosh": np.cosh,
    "tanh": np.tanh,
    "exp": np.exp,
    "log": np.log,
    "log10": np.log10,
    "sqrt": np.sqrt,
    "abs": np.abs,
    "sign": np.sign,
    "floor": np.floor,
    "ceil": np.ceil,
    "round": np.round,
    "deg": _deg,
    "rad": _rad,
    "where": np.where,
    "clip": np.clip,
    "isnan": np.isnan,
    "minimum": np.minimum,
    "maximum": np.maximum,
    "interp": _interp,
    # along the column
    "cumsum": np.nancumsum,
    "diff": np.diff,
    "gradient": np.gradient,
    "smooth": _smooth,
    # one value
    "mean": np.nanmean,
    "median": np.nanmedian,
    "std": np.nanstd,
    "var": np.nanvar,
    "min": np.nanmin,
    "max": np.nanmax,
    "sum": np.nansum,
    "rms": _rms,
    "range": _range,
    "hm0": _hm0,
    "count": lambda x: float(np.count_nonzero(np.isfinite(np.asarray(x, float)))),
    "first": lambda x: float(np.asarray(x, float)[0]),
    "last": lambda x: float(np.asarray(x, float)[-1]),
}
CONSTANTS = {"pi": math.pi, "e": math.e, "g": 9.81, "nan": math.nan, "inf": math.inf}

_OPERATORS = {
    ast.Add: np.add,
    ast.Sub: np.subtract,
    ast.Mult: np.multiply,
    ast.Div: np.true_divide,
    ast.Pow: np.power,
    ast.Mod: np.mod,
    ast.FloorDiv: np.floor_divide,
    ast.BitAnd: np.logical_and,
    ast.BitOr: np.logical_or,
}
_COMPARE = {
    ast.Lt: np.less,
    ast.LtE: np.less_equal,
    ast.Gt: np.greater,
    ast.GtE: np.greater_equal,
    ast.Eq: np.equal,
    ast.NotEq: np.not_equal,
}


def evaluate(expression: str, columns: dict[str, np.ndarray]) -> Any:
    """The value of an expression over the columns (see the module docstring).
    Only arithmetic, comparisons, the columns, :data:`CONSTANTS` and
    :data:`FUNCTIONS`: no attributes, no other names, nothing imported."""
    try:
        tree = ast.parse(expression.strip(), mode="eval")
    except SyntaxError as error:
        raise ValueError(f"cannot read {expression!r}: {error.msg}") from None

    def value(node):
        if isinstance(node, ast.Expression):
            return value(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, int | float):
            return node.value
        if isinstance(node, ast.Constant) and isinstance(node.value, str):
            return node.value  # only for col("name")
        if isinstance(node, ast.Name):
            if node.id in columns:
                return np.asarray(columns[node.id], dtype=np.float64)
            if node.id in CONSTANTS:
                return CONSTANTS[node.id]
            raise KeyError(f"no column {node.id!r}")
        if isinstance(node, ast.BinOp) and type(node.op) in _OPERATORS:
            with np.errstate(divide="ignore", invalid="ignore", over="ignore"):
                return _OPERATORS[type(node.op)](value(node.left), value(node.right))
        if isinstance(node, ast.UnaryOp):
            operand = value(node.operand)
            if isinstance(node.op, ast.USub):
                return np.negative(operand)
            if isinstance(node.op, ast.UAdd):
                return operand
            if isinstance(node.op, ast.Invert | ast.Not):
                return np.logical_not(operand)
        if isinstance(node, ast.Compare):
            left = value(node.left)
            result = None
            for op, comparator in zip(node.ops, node.comparators, strict=True):
                if type(op) not in _COMPARE:
                    break
                right = value(comparator)
                step = _COMPARE[type(op)](left, right)
                result = step if result is None else np.logical_and(result, step)
                left = right
            else:
                return result
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
            if node.keywords:
                raise ValueError("keywords are not allowed in expressions")
            if node.func.id == "col" and len(node.args) == 1:
                name = value(node.args[0])
                if not isinstance(name, str) or name not in columns:
                    raise KeyError(f"no column {name!r}")
                return np.asarray(columns[name], dtype=np.float64)
            function = FUNCTIONS.get(node.func.id)
            if function is None:
                raise ValueError(f"unknown function {node.func.id!r}")
            with np.errstate(divide="ignore", invalid="ignore", over="ignore"):
                return function(*(value(a) for a in node.args))
        if isinstance(node, ast.IfExp):
            return np.where(value(node.test), value(node.body), value(node.orelse))
        raise ValueError(f"not allowed in an expression: {ast.dump(node)[:40]}")

    return value(tree)


# ======================================================================== fits
FIT_MODELS = ("linear", "quadratic", "cubic", "exponential", "power", "logarithmic")


@dataclass
class Fit:
    model: str
    params: dict[str, float]
    r2: float
    equation: str

    def __call__(self, x: np.ndarray) -> np.ndarray:
        x = np.asarray(x, np.float64)
        p = self.params
        if self.model in ("linear", "quadratic", "cubic"):
            coefficients = [p[k] for k in sorted(p, reverse=True)]  # a3, a2, a1, a0
            return np.polyval(coefficients, x)
        if self.model == "exponential":
            return p["a"] * np.exp(p["b"] * x)
        if self.model == "power":
            return p["a"] * np.power(x, p["b"])
        if self.model == "logarithmic":
            return p["a"] + p["b"] * np.log(x)
        raise ValueError(self.model)


def fit(x: np.ndarray, y: np.ndarray, model: str = "linear") -> Fit:
    """A least-squares fit of ``y`` over ``x`` (NaN pairs left out): polynomials
    (linear, quadratic, cubic), exponential a e^(bx), power a x^b, logarithmic
    a + b ln x (the last three fitted in logarithms)."""
    x = np.asarray(x, np.float64)
    y = np.asarray(y, np.float64)
    keep = np.isfinite(x) & np.isfinite(y)
    if model in ("exponential",):
        keep &= y > 0
    if model in ("power",):
        keep &= (x > 0) & (y > 0)
    if model in ("logarithmic",):
        keep &= x > 0
    x, y = x[keep], y[keep]
    degree = {"linear": 1, "quadratic": 2, "cubic": 3}.get(model)
    if x.size < (degree or 1) + 1:
        raise ValueError(f"too few points for a {model} fit")
    if degree is not None:
        coefficients = np.polyfit(x, y, degree)
        params = {f"a{degree - k}": float(c) for k, c in enumerate(coefficients)}
        terms = []
        for k, c in enumerate(coefficients):
            power = degree - k
            terms.append(f"{c:.4g}" + ("" if power == 0 else " x" if power == 1 else f" x^{power}"))
        equation = "y = " + " + ".join(terms).replace("+ -", "- ")
    elif model == "exponential":
        b, log_a = np.polyfit(x, np.log(y), 1)
        params = {"a": float(np.exp(log_a)), "b": float(b)}
        equation = f"y = {params['a']:.4g} exp({params['b']:.4g} x)"
    elif model == "power":
        b, log_a = np.polyfit(np.log(x), np.log(y), 1)
        params = {"a": float(np.exp(log_a)), "b": float(b)}
        equation = f"y = {params['a']:.4g} x^{params['b']:.4g}"
    elif model == "logarithmic":
        b, a = np.polyfit(np.log(x), y, 1)
        params = {"a": float(a), "b": float(b)}
        equation = f"y = {params['a']:.4g} + {params['b']:.4g} ln x"
    else:
        raise ValueError(f"unknown fit {model!r}; one of {', '.join(FIT_MODELS)}")
    result = Fit(model, params, float("nan"), equation)
    residual = y - result(x)
    spread = np.sum((y - y.mean()) ** 2)
    result.r2 = float(1.0 - np.sum(residual**2) / spread) if spread > 0 else float("nan")
    return result


# ======================================================================== where from
@dataclass
class TableSpec:
    """Where a table comes from and its computed columns (kept in a figure)."""

    name: str
    source: str  # "file:<path>", "series:<key>", "profile:<key>@<t>", "collect:"
    formulas: dict[str, str] = field(default_factory=dict)
    #: collect: one row per open case, these values (name -> "key: expression")
    collect: dict[str, str] = field(default_factory=dict)
    start: float | None = None  # collect: the time window
    end: float | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> TableSpec:
        return cls(**{k: v for k, v in data.items() if k in cls.__dataclass_fields__})

    @property
    def path(self) -> Path | None:
        return Path(self.source[5:]) if self.source.startswith("file:") else None

    def load(self, session=None, base: Path | None = None) -> DataTable:
        """The table, its formulas computed. ``base``: where a relative file is."""
        kind, _, rest = self.source.partition(":")
        if kind == "file":
            path = Path(rest).expanduser()
            if not path.is_absolute() and base is not None:
                path = Path(base) / path
            table = read_table(path, self.name)
        elif kind == "series":
            table = series_table(session, rest, self.name)
        elif kind == "profile":
            key, _, when = rest.partition("@")
            table = profile_table(session, key, float(when or 0.0), self.name)
        elif kind == "collect":
            table = collect(session, self.collect, self.start, self.end, self.name)
        else:
            raise ValueError(f"{self.name}: unknown source {self.source!r}")
        table.name = self.name
        table.source = self.source
        if self.formulas:
            table.add_formulas(dict(self.formulas))
        return table

    def stamp(self, base: Path | None = None) -> float:
        """When the file of the table changed (0 for other sources): to notice an
        update and draw again."""
        path = self.path
        if path is None:
            return 0.0
        if not path.is_absolute() and base is not None:
            path = Path(base) / path
        try:
            return path.stat().st_mtime
        except OSError:
            return 0.0


def series_table(session, key: str, name: str = "") -> DataTable:
    """A time-series file of the first open case that has it, as a table: "time"
    and its columns."""
    for item in session.scene:
        if item.kind != "case":
            continue
        case = session.data.case(item.id)
        if case is None or key not in case.series:
            continue
        series = case.series[key].read()
        table = DataTable(name or key, source=f"series:{key}")
        table.columns["time"] = np.asarray(series.time, np.float64)
        table.units["time"] = "s"
        for column, values in series.columns.items():
            table.columns[column] = np.asarray(values, np.float64)
            if series.units.get(column):
                table.units[column] = series.units[column]
        return table
    raise KeyError(f"no open case has the time series {key!r}")


def profile_table(session, key: str, time: float, name: str = "") -> DataTable:
    """A free surface line set at a time: "x" and a column per line ("eta1", "eta2",
    ...; the wave theory as "theory1", ...)."""
    for item in session.scene:
        if item.kind != "case":
            continue
        case = session.data.case(item.id)
        if case is None or key not in case.profiles:
            continue
        profile = case.profiles[key].at(float(time))
        table = DataTable(name or f"{key}@{time:g}", source=f"profile:{key}@{time:g}")
        for number in sorted(profile.lines):
            line = profile.lines[number]
            if "x" not in table.columns:
                table.columns["x"] = np.asarray(line.x, np.float64)
                table.units["x"] = "m"
            if len(line.x) != len(table.columns["x"]):
                continue
            table.columns[f"eta{number}"] = np.asarray(line.eta, np.float64)
            table.units[f"eta{number}"] = "m"
            if line.theory is not None:
                table.columns[f"theory{number}"] = np.asarray(line.theory, np.float64)
        return table
    raise KeyError(f"no open case has the free surface lines {key!r}")


def collect(
    session,
    values: dict[str, str],
    start: float | None = None,
    end: float | None = None,
    name: str = "runs",
) -> DataTable:
    """One row per open case (a parameter study, a grid convergence): for each value
    ``"key: expression"`` the expression over that case's time-series file ``key``
    in the time window, e.g. ``{"Fz max": "nhflow_6dof_forces_0: max(Fz)",
    "Hm0": "nhflow_wsf_hg: hm0(P3)"}``. Columns: "case" (1, 2, ...) and the values;
    a case without the file gets NaN. The cases' names are the table's ``labels``."""
    cases = [item for item in session.scene if item.kind == "case"]
    table = DataTable(name, source="collect:")
    table.columns["case"] = np.arange(1, len(cases) + 1, dtype=np.float64)
    table.labels = [item.name for item in cases]
    for column, recipe in values.items():
        key, _, expression = recipe.partition(":")
        key, expression = key.strip(), expression.strip()
        found = np.full(len(cases), np.nan)
        for row, item in enumerate(cases):
            case = session.data.case(item.id)
            if case is None or key not in case.series:
                continue
            series = case.series[key].read()
            time = np.asarray(series.time, np.float64)
            keep = np.ones(time.size, bool)
            if start is not None:
                keep &= time >= start
            if end is not None:
                keep &= time <= end
            columns = {"time": time[keep]}
            columns.update({c: np.asarray(v, np.float64)[keep] for c, v in series.columns.items()})
            try:
                result = np.asarray(evaluate(expression, columns), np.float64)
            except (ValueError, KeyError, TypeError) as error:
                table.errors[column] = str(error)
                continue
            found[row] = float(result) if result.ndim == 0 else float(np.nanmean(result))
        table.columns[column] = found
    return table
